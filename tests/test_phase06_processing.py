"""Synthetic Phase06 A tests + real isolated PDF/DOCX/Persian OCR B fixtures."""

from __future__ import annotations

import asyncio
import hashlib
import io
import shutil
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.core.services.ingestion.ingestion_worker_service import IngestionWorkerService
from src.core.services.ingestion.upload_validation import UploadRejected, validate_content
from src.core.services.persian_text.persian_processing import (
    ExtractedPage,
    ProcessingRejected,
    chunk_pages,
    normalize_persian,
)
from src.infrastructure.providers.document.document_extractor import DocumentExtractor
from src.infrastructure.utils.config_reader import ProcessingSettings

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_loss_aware_persian_arabic_mixed_normalization():
    raw = "علي كاظمي\u00a0در ماده ١٢ و سال ۱۴۰۵ گفت: آیه \u200c۱۲"
    expected = "علی کاظمی در ماده 12 و سال 1405 گفت: آیه \u200c12"
    assert normalize_persian(raw) == expected
    assert "علي" in raw and "علي" not in normalize_persian(raw)
    assert normalize_persian("Hello GPT-6 | سلام ٣") == "Hello GPT-6 | سلام 3"
    assert normalize_persian("هدى") == "هدى"  # Arabic alef maqsura is not Persian yeh.


def test_stable_chunk_ids_and_true_raw_spans():
    version_id = uuid4()
    raw = "ماده ۱۲\nشرح قرارداد و آثار آن.\n\nسؤال: مسئول کیست؟\nپاسخ: مدیر واحد."
    pages = [ExtractedPage(2, raw, "digital_pdf")]
    first = chunk_pages(version_id, pages, max_chars=256)
    second = chunk_pages(version_id, pages, max_chars=256)
    assert first == second
    assert first[0].page == 2
    assert all(raw[item.start : item.end] == item.original for item in first)
    assert all(item.checksum == hashlib.sha256(item.canonical.encode()).hexdigest() for item in first)
    assert all(item.id.version == 5 for item in first)
    assert first[0].section is not None
    assert first[0].canonical.startswith("ماده 12")
    assert chunk_pages(uuid4(), pages, max_chars=256)[0].id != first[0].id


def test_no_fabricated_citation_or_empty_chunk():
    with pytest.raises(ProcessingRejected, match="no_extractable_text"):
        chunk_pages(uuid4(), [ExtractedPage(None, "   \n\n", "utf8_text")])
    with pytest.raises(ProcessingRejected, match="invalid_page_number"):
        chunk_pages(uuid4(), [ExtractedPage(0, "Hello", "digital_pdf")])
    with pytest.raises(ValueError, match="Invalid chunk"):
        chunk_pages(uuid4(), [ExtractedPage(None, "Hello", "utf8_text")], max_chars=200)


def test_long_document_is_bounded_and_deterministic():
    raw = "ماده ۱.\n" + ("این بند شامل عبارت فارسی و استناد صحیح است. " * 90)
    result = chunk_pages(uuid4(), [ExtractedPage(None, raw, "utf8_text")], max_chars=256)
    assert len(result) >= 4
    assert all(0 < item.end - item.start <= 256 for item in result)
    assert all(raw[item.start : item.end] == item.original for item in result)


def _digital_pdf() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 20 Tf 72 720 Td (Hello digital PDF evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(contents)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_real_digital_pdf_extraction_page_and_span():
    data = _digital_pdf()
    validate_content(data, "application/pdf")
    pages = asyncio.run(DocumentExtractor().extract_async(data, "application/pdf"))
    assert len(pages) == 1
    assert pages[0].number == 1
    assert pages[0].method == "digital_pdf"
    assert "digital PDF evidence" in pages[0].text
    pieces = chunk_pages(uuid4(), pages)
    assert pieces[0].original == pages[0].text[pieces[0].start : pieces[0].end]


def test_custom_pdf_page_limit_reaches_real_extractor_child():
    writer = PdfWriter()
    page = PdfReader(io.BytesIO(_digital_pdf())).pages[0]
    writer.add_page(page)
    writer.add_page(page)
    output = io.BytesIO()
    writer.write(output)
    data = output.getvalue()

    denied_limits = ProcessingSettings(max_pdf_pages=1, max_ocr_pages=1)
    with pytest.raises(ProcessingRejected, match="extraction_failed"):
        asyncio.run(DocumentExtractor().extract_async(data, "application/pdf", limits=denied_limits))

    allowed_limits = ProcessingSettings(max_pdf_pages=2, max_ocr_pages=1)
    pages = asyncio.run(DocumentExtractor().extract_async(data, "application/pdf", limits=allowed_limits))
    assert [item.number for item in pages] == [1, 2]


def test_scanned_worker_forwards_resource_limits_after_malware_gate():
    data = b"Synthetic safe content for config propagation"
    version_id = uuid4()
    job_id = uuid4()
    checksum = hashlib.sha256(data).hexdigest()
    processing = ProcessingSettings(max_pdf_pages=1, max_ocr_pages=1, max_chunk_chars=256)
    events = []

    class Repository:
        async def claim_async(self, worker_id, scope, lease_seconds):
            assert scope == "internal"
            return job_id, "synthetic-lease", "synthetic-key", checksum, len(data), version_id, "text/plain"

        async def finish_async(self, received_id, token, ok, code, retriable, *, chunks):
            assert received_id == job_id and token == "synthetic-lease"
            assert ok and code is None and not retriable
            assert len(chunks) == 1 and chunks[0].canonical == "سلام"
            events.append("finish")
            return True

    class Storage:
        async def get_async(self, key):
            events.append("storage")
            return data

    class Scanner:
        async def scan_async(self, contents):
            assert contents == data
            events.append("scan")

    class Extractor:
        async def extract_async(self, contents, mime, *, limits):
            assert contents == data and mime == "text/plain"
            assert limits is processing
            events.append("extract")
            return [ExtractedPage(None, "سلام", "utf8_text")]

    worker = object.__new__(IngestionWorkerService)
    worker._repo = Repository()
    worker._storage = Storage()
    worker._scanner = Scanner()
    worker._extractor = Extractor()
    worker._cfg = SimpleNamespace(lease_seconds=90)
    worker._scope = "internal"
    worker._processing = processing

    assert asyncio.run(worker.run_once_async("config-test"))
    assert events == ["storage", "scan", "extract", "finish"]


def test_real_docx_paragraphs_and_table_extraction():
    document = Document()
    document.add_heading("بخش اول", level=1)
    document.add_paragraph("متن فارسی و Arabic")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "عنوان", "پاسخ"
    table.cell(1, 0).text, table.cell(1, 1).text = "سؤال", "گزینه"
    output = io.BytesIO()
    document.save(output)
    data = output.getvalue()
    validate_content(data, DOCX_MIME)
    pages = asyncio.run(DocumentExtractor().extract_async(data, DOCX_MIME))
    assert len(pages) == 1
    assert pages[0].number is None  # Word has no defensible fixed PDF page coordinates.
    assert "# بخش اول" in pages[0].text  # OOXML Heading 1 remains a true structural boundary.
    assert "عنوان | پاسخ" in pages[0].text
    assert "سؤال | گزینه" in pages[0].text
    assert all(item.page is None for item in chunk_pages(uuid4(), pages))


def test_encrypted_pdf_is_never_extracted_or_approved():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("synthetic-secret")
    output = io.BytesIO()
    writer.write(output)
    data = output.getvalue()
    with pytest.raises(UploadRejected):
        validate_content(data, "application/pdf")
    with pytest.raises(ProcessingRejected, match="extraction_failed"):
        asyncio.run(DocumentExtractor().extract_async(data, "application/pdf"))


def test_malformed_word_and_pdf_fail_closed():
    with pytest.raises(UploadRejected):
        validate_content(b"PK\x03\x04broken", DOCX_MIME)
    with pytest.raises(UploadRejected):
        validate_content(b"%PDF-missing-trailer", "application/pdf")
    with pytest.raises(ProcessingRejected):
        asyncio.run(DocumentExtractor().extract_async(b"not a PDF", "application/pdf"))


def _character_error_rate(actual: str, expected: str) -> float:
    """Levenshtein distance divided by ground-truth codepoint count."""
    previous = list(range(len(expected) + 1))
    for i, character in enumerate(actual, 1):
        row = [i]
        for j, reference in enumerate(expected, 1):
            row.append(
                min(
                    previous[j] + 1,
                    row[j - 1] + 1,
                    previous[j - 1] + (character != reference),
                )
            )
        previous = row
    return previous[-1] / len(expected)


@pytest.mark.skipif(
    not shutil.which("pdftoppm") or not shutil.which("tesseract"),
    reason="Real Poppler and Tesseract OCR binaries not provisioned",
)
@pytest.mark.parametrize("expected, language", [("سلام دنیا", "fa"), ("مرحبا بالعالم", "ar")])
def test_real_scanned_persian_pdf_ocr(expected: str, language: str):
    # PIL raster-only PDF: no hidden text layer and no simulated OCR response.
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 75)
    image = Image.new("RGB", (1300, 240), "white")
    ImageDraw.Draw(image).text((1240, 48), expected, font=font, fill="black", anchor="ra", direction="rtl")
    pdf = io.BytesIO()
    image.save(pdf, format="PDF", resolution=180)
    data = pdf.getvalue()
    validate_content(data, "application/pdf")
    pages = asyncio.run(DocumentExtractor().extract_async(data, "application/pdf"))
    assert pages[0].method == "tesseract_ocr"
    assert pages[0].number == 1
    observed = normalize_persian(pages[0].text)
    cer = _character_error_rate(observed, expected)
    print(f"Phase06 scanned-PDF {language} OCR synthetic CER={cer:.3f}")
    assert cer <= 0.20, (observed, expected, cer)
    assert chunk_pages(uuid4(), pages)[0].method == "tesseract_ocr"


@pytest.mark.skipif(
    not shutil.which("pdftoppm") or not shutil.which("tesseract"),
    reason="Real Poppler and Tesseract OCR binaries not provisioned",
)
def test_mixed_digital_and_raster_pdf_uses_real_ocr():
    image = Image.new("RGB", (1300, 1800), "white")
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 95)
    ImageDraw.Draw(image).text((100, 900), "RASTER KNOWLEDGE", font=font, fill="black")
    raster_pdf = io.BytesIO()
    image.save(raster_pdf, format="PDF", resolution=180)
    image_page = PdfReader(io.BytesIO(raster_pdf.getvalue())).pages[0]
    text_page = PdfReader(io.BytesIO(_digital_pdf())).pages[0]
    image_page.merge_page(text_page)
    writer = PdfWriter()
    writer.add_page(image_page)
    output = io.BytesIO()
    writer.write(output)

    pages = asyncio.run(DocumentExtractor().extract_async(output.getvalue(), "application/pdf"))
    assert len(pages) == 1 and pages[0].number == 1
    assert pages[0].method == "hybrid_pdf_ocr"
    assert "Hello digital PDF evidence" in pages[0].text
    assert "RASTER KNOWLEDGE" in pages[0].text
    assert "tesseract" in pages[0].engine_version
    chunks = chunk_pages(uuid4(), pages)
    assert all(pages[0].text[item.start : item.end] == item.original for item in chunks)


@pytest.mark.skipif(
    not shutil.which("pdftoppm") or not shutil.which("tesseract"),
    reason="Real Poppler and Tesseract OCR binaries not provisioned",
)
def test_real_blank_scanned_pdf_is_rejected_before_approval():
    image = Image.new("RGB", (800, 200), "white")
    pdf = io.BytesIO()
    image.save(pdf, format="PDF", resolution=180)
    data = pdf.getvalue()
    validate_content(data, "application/pdf")
    with pytest.raises(ProcessingRejected, match="extraction_failed"):
        asyncio.run(DocumentExtractor().extract_async(data, "application/pdf"))


def test_extractor_timeout_terminates_child(monkeypatch):
    from src.infrastructure.providers.document import document_extractor

    class HangingChild:
        def __init__(self):
            self.killed = False

        async def communicate(self, contents):
            raise TimeoutError("simulated hung subprocess")

        def kill(self):
            self.killed = True

        async def wait(self):
            return 1

    child = HangingChild()

    async def create_child(*args, **kwargs):
        return child

    monkeypatch.setattr(document_extractor.asyncio, "create_subprocess_exec", create_child)
    with pytest.raises(ProcessingRejected, match="extraction_timeout"):
        asyncio.run(DocumentExtractor().extract_async(b"safe", "text/plain"))
    assert child.killed
