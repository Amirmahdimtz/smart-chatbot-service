"""Phase 05 Group B real PostgreSQL + S3-compatible + clamd integration.

Runs ONLY with disposable local synthetic services explicitly enabled by CI.
"""

from __future__ import annotations

import asyncio
import base64
import os
import subprocess
import sys
import time
from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import UUID, uuid4

import httpx
import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from fastapi.testclient import TestClient

from src.core.services.ingestion.ingestion_worker_service import IngestionWorkerService
from src.host.app import create_app
from src.infrastructure.context.sql_db.psql_dbcontext import PsqlDbContext
from src.infrastructure.di.bootstrap import bootstrap_di
from src.infrastructure.models.ingestion import IngestionJob
from src.infrastructure.providers.scanner.clamav_scanner import ClamAvScanner, MalwareDetected
from src.infrastructure.providers.storage.quarantine_storage import QuarantineStorage
from src.infrastructure.repositories.ingestion.ingestion_repository import IngestionRepository
from src.infrastructure.security.token_validator import TokenValidator

EICAR = b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"


def _enabled():
    if os.environ.get("SCB_PHASE05_TEST_ENABLED") != "true":
        pytest.skip("Actual isolated phase05 dependencies not configured")
    assert os.environ["SCB_ENV"] == "development"
    assert os.environ["SCB_ZONE"] == "internal"
    assert os.environ["SCB_DATABASE_ENABLED"] == "true"
    assert os.environ["SCB_INGESTION_ENABLED"] == "true"
    assert os.environ["SCB_DATABASE_URL"].endswith("/scb_phase05_test")
    assert os.environ["SCB_INGESTION_S3_URL"] == "http://127.0.0.1:9000"
    assert os.environ["SCB_INGESTION_CLAMD_HOST"] == "127.0.0.1"


def _auth():
    private = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    public = private.public_key().public_numbers()

    def b64(n):
        return base64.urlsafe_b64encode(n.to_bytes((n.bit_length() + 7) // 8, "big")).rstrip(b"=").decode()

    jwks = {
        "keys": [
            {"kty": "RSA", "use": "sig", "alg": "RS256", "kid": "synthetic", "n": b64(public.n), "e": b64(public.e)}
        ]
    }

    def token(subject, permissions=None, audience="synthetic-internal"):
        now = datetime.now(UTC)
        return jwt.encode(
            {
                "iss": "https://identity.example.test/",
                "aud": audience,
                "sub": subject,
                "iat": now,
                "nbf": now,
                "exp": now + timedelta(minutes=10),
                "permissions": permissions or [],
            },
            private,
            algorithm="RS256",
            headers={"kid": "synthetic"},
        )

    return jwks, token


@pytest.fixture
def client(monkeypatch):
    _enabled()
    monkeypatch.setenv("SCB_JWKS_URL", "https://identity.example.test/jwks")
    monkeypatch.setenv("SCB_ISSUER", "https://identity.example.test/")
    monkeypatch.setenv("SCB_AUDIENCE", "synthetic-internal")
    app = create_app()
    jwks, token = _auth()
    validator = app.state.container.resolve(TokenValidator)

    async def load_keys(request):
        return httpx.Response(200, json=jwks)

    validator._client = httpx.AsyncClient(transport=httpx.MockTransport(load_keys))
    with TestClient(app) as http:
        yield http, token, app


def test_live_upload_to_durable_job_and_independent_worker(client):
    http, token, app = client
    owner = "synthetic-" + uuid4().hex
    bearer = {"Authorization": "Bearer " + token(owner, ["knowledge:test:ingest"])}
    response = http.post("/api/v1/knowledge/sources", headers=bearer, json={"domain": "research"})
    assert response.status_code == 201, response.text
    source = response.json()["id"]
    payload = ("Synthetic Persian knowledge: " + "سلام").encode()
    idem = uuid4().hex
    uploaded = http.post(
        f"/api/v1/knowledge/sources/{source}/documents",
        headers={**bearer, "Idempotency-Key": idem},
        files={"file": ("paper.txt", payload, "text/plain")},
    )
    assert uploaded.status_code == 202, uploaded.text
    job = uploaded.json()
    assert job["status"] == "queued"
    assert http.get(f"/api/v1/knowledge/jobs/{job['id']}", headers=bearer).status_code == 200
    assert (
        http.post(
            f"/api/v1/knowledge/sources/{source}/documents",
            headers={**bearer, "Idempotency-Key": idem},
            files={"file": ("paper.txt", payload, "text/plain")},
        ).json()["id"]
        == job["id"]
    )

    # A second idempotency key for identical bytes must be rejected instead of
    # silently aliasing to a job without remembering that alternative key.
    collision = http.post(
        f"/api/v1/knowledge/sources/{source}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex},
        files={"file": ("paper.txt", payload, "text/plain")},
    )
    assert collision.status_code == 409
    assert collision.json() == {"error": "idempotency_conflict"}

    outsider = {"Authorization": "Bearer " + token("synthetic-foreign", ["knowledge:test:ingest"])}
    assert (
        http.post(
            f"/api/v1/knowledge/sources/{source}/documents",
            headers={**outsider, "Idempotency-Key": uuid4().hex},
            files={"file": ("paper.txt", b"foreign", "text/plain")},
        ).status_code
        == 403
    )

    for filename, body, mime, expected in (
        ("../traversal.txt", b"safe", "text/plain", 422),
        ("not-a-pdf.pdf", b"%PDF-broken", "application/pdf", 422),
        ("archive.zip", b"PK\\x03\\x04foo", "application/zip", 415),
        ("huge.txt", b"a" * 1048577, "text/plain", 413),
    ):
        rejected = http.post(
            f"/api/v1/knowledge/sources/{source}/documents",
            headers={**bearer, "Idempotency-Key": uuid4().hex},
            files={"file": (filename, body, mime)},
        )
        assert rejected.status_code == expected, (filename, rejected.status_code, rejected.text)
    assert (
        http.post(
            f"/api/v1/knowledge/sources/{source}/documents",
            headers={**bearer, "Idempotency-Key": uuid4().hex, "x-user-id": "synthetic-victim"},
            files={"file": ("ok.txt", b"safe", "text/plain")},
        ).status_code
        == 400
    )
    malformed = http.post(
        f"/api/v1/knowledge/sources/{source}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex, "content-type": "multipart/form-data; boundary=broken"},
        content=b"invalid multipart payload",
    )
    assert malformed.status_code in {400, 422}

    # Real OS process; no request-loop background queue. Poll the actual persisted job.
    process = subprocess.Popen(
        [sys.executable, "-m", "src.host.ingestion_worker"],
        env=os.environ.copy(),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            fetched = http.get(f"/api/v1/knowledge/jobs/{job['id']}", headers=bearer)
            if fetched.status_code == 200 and fetched.json()["status"] == "completed":
                break
            assert process.poll() is None, "Worker exited unexpectedly"
            time.sleep(0.3)
        else:
            pytest.fail("Worker did not persist completed job")
        version_id = job["version_id"]
        version = http.get(f"/api/v1/knowledge/versions/{version_id}", headers=bearer)
        assert version.status_code == 200
        assert version.json()["state"] == "awaiting_approval"
        assert "published" not in version.text

        # No other owner / no other audience can observe the job.
        stranger = {"Authorization": "Bearer " + token("synthetic-other", ["knowledge:test:ingest"])}
        assert http.get(f"/api/v1/knowledge/jobs/{job['id']}", headers=stranger).status_code == 404
        assert (
            http.get(
                f"/api/v1/knowledge/jobs/{job['id']}", headers={"Authorization": "Bearer " + token(owner, [])}
            ).status_code
            == 403
        )
        assert (
            http.get(
                f"/api/v1/knowledge/jobs/{job['id']}",
                headers={"Authorization": "Bearer " + token(owner, ["knowledge:test:ingest"], audience="other")},
            ).status_code
            == 401
        )
        assert http.get("/api/v1/knowledge/jobs/" + job["id"]).status_code == 401
    finally:
        process.terminate()
        process.wait(timeout=10)


def test_actual_s3_gateway_rejects_invalid_and_unsigned_credentials():
    """A real S3 gateway must fail closed, unlike an unauthenticated mock."""
    _enabled()
    import boto3
    from botocore import UNSIGNED
    from botocore.config import Config as BotoConfig
    from botocore.exceptions import ClientError

    endpoint = os.environ["SCB_INGESTION_S3_URL"]
    bucket = os.environ["SCB_INGESTION_S3_BUCKET"]
    # Even a healthy S3 service must not grant reads to an unknown key.
    bad = boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id="invalid-ci-key",
        aws_secret_access_key="invalid-ci-secret",
        config=BotoConfig(s3={"addressing_style": "path"}, connect_timeout=3, read_timeout=3),
    )
    try:
        with pytest.raises(ClientError) as invalid:
            bad.head_bucket(Bucket=bucket)
        assert invalid.value.response["ResponseMetadata"]["HTTPStatusCode"] in {401, 403}
    finally:
        bad.close()

    anonymous = boto3.client(
        "s3",
        endpoint_url=endpoint,
        config=BotoConfig(
            signature_version=UNSIGNED, s3={"addressing_style": "path"}, connect_timeout=3, read_timeout=3
        ),
    )
    try:
        with pytest.raises(ClientError) as unsigned:
            anonymous.get_object(Bucket=bucket, Key="quarantine/internal/" + str(uuid4()))
        assert unsigned.value.response["ResponseMetadata"]["HTTPStatusCode"] in {401, 403}
    finally:
        anonymous.close()


def test_real_scanner_and_quarantine_storage():
    _enabled()
    from src.infrastructure.providers.storage.quarantine_storage import StorageUnavailable

    container = bootstrap_di()
    scanner = container.resolve(ClamAvScanner)
    storage = container.resolve(QuarantineStorage)

    async def exercise():
        await storage.start_async()
        try:
            clean = b"synthetic dataset"
            key = f"quarantine/internal/{uuid4()}"
            import hashlib

            checksum = hashlib.sha256(clean).hexdigest()
            await storage.put_async(key, clean, checksum)
            await storage.put_async(key, clean, checksum)  # Lost ACK, safe identical replay.
            assert await storage.get_async(key) == clean
            with pytest.raises(StorageUnavailable):
                await storage.put_async(key, b"modified", hashlib.sha256(b"modified").hexdigest())
            assert await storage.get_async(key) == clean
            await scanner.scan_async(clean)
            with pytest.raises(MalwareDetected):
                await scanner.scan_async(EICAR)
        finally:
            await storage.close_async()

    asyncio.run(exercise())


def test_durable_concurrent_claim_fencing_retry_and_cancellation():
    _enabled()
    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    repo = container.resolve(IngestionRepository)
    storage = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def exercise():
        await db.start_async()
        await storage.start_async()
        try:
            source = await repo.create_source_async("synthetic-" + uuid4().hex, "internal", "research")
            import hashlib

            content = ("synthetic-" + uuid4().hex).encode()
            checksum = hashlib.sha256(content).hexdigest()
            job, version, created = await repo.stage_async(
                source.id,
                source.owner_principal_id,
                "internal",
                "test.txt",
                "text/plain",
                len(content),
                checksum,
                uuid4().hex,
                2,
            )
            assert created
            await storage.put_async(version.object_key, content, checksum)
            assert await repo.mark_queued_async(job.id, source.owner_principal_id, "internal")
            claimed = await asyncio.gather(
                repo.claim_async("workerA", "internal", 45),
                repo.claim_async("workerB", "internal", 45),
            )
            valid = [item for item in claimed if item and item[0] == job.id]
            assert len(valid) == 1
            first = valid[0]
            assert await repo.heartbeat_async(job.id, first[1], 45)
            assert not await repo.heartbeat_async(job.id, "stale-token", 45)
            assert await repo.finish_async(job.id, first[1], False, "dependency_unavailable", True)
            assert not await repo.finish_async(job.id, first[1], True)
            async with db.transaction_async() as session:
                j = await session.get(IngestionJob, job.id)
                assert j.status == "retry_wait"
                j.next_retry_at = datetime.now(UTC) - timedelta(seconds=1)
            second = await repo.claim_async("workerC", "internal", 45)
            assert second and second[0] == job.id
            assert second[1] != first[1]
            assert await repo.finish_async(job.id, second[1], False, "dependency_unavailable", True)
            after = await repo.get_job_owned_async(job.id, source.owner_principal_id, "internal")
            assert after[0].status == "dead"
            assert after[0].attempts == 2
            assert after[1].state == "failed"
            another = ("synthetic-" + uuid4().hex).encode()
            other_hash = hashlib.sha256(another).hexdigest()
            next_job, next_ver, _ = await repo.stage_async(
                source.id,
                source.owner_principal_id,
                "internal",
                "another.txt",
                "text/plain",
                len(another),
                other_hash,
                uuid4().hex,
                2,
            )
            await storage.put_async(next_ver.object_key, another, other_hash)
            assert await repo.mark_queued_async(next_job.id, source.owner_principal_id, "internal")
            assert await repo.cancel_async(next_job.id, source.owner_principal_id, "internal")
            assert not await repo.cancel_async(next_job.id, source.owner_principal_id, "internal")
            assert (await repo.get_job_owned_async(next_job.id, source.owner_principal_id, "internal"))[
                0
            ].status == "cancelled"
            assert await worker.run_once_async("worker-after-cancel") is False
        finally:
            await storage.close_async()
            await db.close_async()

    asyncio.run(exercise())


def test_real_post_put_crash_recovery_missing_staging_and_lease_expiry():
    """Reconstruct committed staging after crash; never silently lose an orphaned intent."""
    _enabled()
    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    repo = container.resolve(IngestionRepository)
    store = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def exercise():
        import hashlib

        await db.start_async()
        await store.start_async()
        try:
            actor = "synthetic-" + uuid4().hex
            source = await repo.create_source_async(actor, "internal", "research")

            # Simulates the API process dying after S3 PUT, before durable queue ACK.
            content = ("crash-after-put-" + uuid4().hex).encode()
            checksum = hashlib.sha256(content).hexdigest()
            key = uuid4().hex
            job, version, fresh = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "crash.txt",
                "text/plain",
                len(content),
                checksum,
                key,
                2,
            )
            assert fresh and job.status == "staging"
            await store.put_async(version.object_key, content, checksum)
            async with db.transaction_async() as session:
                j = await session.get(IngestionJob, job.id)
                j.created_at = datetime.now(UTC) - timedelta(seconds=35)
            assert await worker.run_once_async("recovery-worker")
            recovered = await repo.get_job_owned_async(job.id, actor, "internal")
            assert recovered[0].status == "queued"
            assert await worker.run_once_async("recovery-worker")
            completed = await repo.get_job_owned_async(job.id, actor, "internal")
            assert completed[0].status == "completed"
            assert completed[1].state == "awaiting_approval"

            # Simulates a crash after committing the job but before writing the S3 object.
            missing = ("not-uploaded-" + uuid4().hex).encode()
            missing_checksum = hashlib.sha256(missing).hexdigest()
            idem = uuid4().hex
            lost, lost_version, fresh = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "missing.txt",
                "text/plain",
                len(missing),
                missing_checksum,
                idem,
                2,
            )
            assert fresh
            async with db.transaction_async() as session:
                j = await session.get(IngestionJob, lost.id)
                j.created_at = datetime.now(UTC) - timedelta(seconds=360)
            assert await worker.run_once_async("recovery-worker")
            old = await repo.get_job_owned_async(lost.id, actor, "internal")
            assert old[0].status == "dead"
            assert old[0].error_code == "staging_abandoned"
            # Same idempotency key and checksum may safely resume a pre-processing failed upload.
            resumed, resumable_version, fresh = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "missing.txt",
                "text/plain",
                len(missing),
                missing_checksum,
                idem,
                2,
            )
            assert fresh and resumed.id == lost.id and resumable_version.id == lost_version.id
            await store.put_async(resumable_version.object_key, missing, missing_checksum)
            assert await repo.mark_queued_async(resumed.id, actor, "internal")
            assert await worker.run_once_async("recovery-worker")
            assert (await repo.get_job_owned_async(resumed.id, actor, "internal"))[0].status == "completed"

            # A crashed final-attempt worker cannot be claimed for an unbounded extra attempt.
            third = ("lease-expiration-" + uuid4().hex).encode()
            third_checksum = hashlib.sha256(third).hexdigest()
            final_job, final_version, _ = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "lease.txt",
                "text/plain",
                len(third),
                third_checksum,
                uuid4().hex,
                1,
            )
            await store.put_async(final_version.object_key, third, third_checksum)
            assert await repo.mark_queued_async(final_job.id, actor, "internal")
            claim = await repo.claim_async("crashed-worker", "internal", 45)
            assert claim[0] == final_job.id
            async with db.transaction_async() as session:
                j = await session.get(IngestionJob, final_job.id)
                j.lease_expires_at = datetime.now(UTC) - timedelta(seconds=2)
            assert await worker.run_once_async("recovery-worker")
            assert not await repo.finish_async(final_job.id, claim[1], True)
            final = await repo.get_job_owned_async(final_job.id, actor, "internal")
            assert final[0].status == "dead"
            assert final[0].error_code == "lease_expired"
            assert final[1].state == "failed"
        finally:
            await store.close_async()
            await db.close_async()

    asyncio.run(exercise())


def test_real_malware_corrupted_object_and_scanner_network_failure():
    """Real S3 bytes + real ClamAV verdict; none may reach awaiting_approval."""
    _enabled()
    import hashlib
    import socket
    from types import SimpleNamespace

    from src.infrastructure.providers.scanner.clamav_scanner import ScanUnavailable

    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    repo = container.resolve(IngestionRepository)
    store = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def exercise():
        await db.start_async()
        await store.start_async()
        try:
            actor = "synthetic-" + uuid4().hex
            source = await repo.create_source_async(actor, "internal", "research")
            malware_hash = hashlib.sha256(EICAR).hexdigest()
            infected, infected_version, created = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "eicar.txt",
                "text/plain",
                len(EICAR),
                malware_hash,
                uuid4().hex,
                1,
            )
            assert created
            await store.put_async(infected_version.object_key, EICAR, malware_hash)
            assert await repo.mark_queued_async(infected.id, actor, "internal")
            assert await worker.run_once_async("scanner-real")
            scanned = await repo.get_job_owned_async(infected.id, actor, "internal")
            assert scanned[0].status == "dead"
            assert scanned[0].error_code == "malware_detected"
            assert scanned[1].state == "failed"

            original = ("expected-" + uuid4().hex).encode()
            changed = ("modified-" + uuid4().hex).encode()
            queued, version, _ = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "changed.txt",
                "text/plain",
                len(original),
                hashlib.sha256(original).hexdigest(),
                uuid4().hex,
                1,
            )
            await store.put_async(version.object_key, original, hashlib.sha256(original).hexdigest())
            assert await repo.mark_queued_async(queued.id, actor, "internal")
            # Simulate a malicious or accidental object overwrite after HEAD verification.
            # Bypass the production adapter to model an out-of-band actor with
            # direct bucket privileges. The safe adapter itself forbids overwrites.
            await asyncio.to_thread(
                store._client.put_object,
                Bucket=store._cfg.s3_bucket,
                Key=version.object_key,
                Body=changed,
                Metadata={"sha256": hashlib.sha256(changed).hexdigest()},
            )
            assert await worker.run_once_async("tamper-detection")
            tampered = await repo.get_job_owned_async(queued.id, actor, "internal")
            assert tampered[0].status == "dead"
            assert tampered[0].error_code == "integrity_mismatch"
            assert tampered[1].state == "failed"

            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
                probe.bind(("127.0.0.1", 0))
                unused_port = probe.getsockname()[1]
            cfg = SimpleNamespace(
                settings=SimpleNamespace(
                    ingestion=SimpleNamespace(
                        enabled=True,
                        max_file_bytes=1048576,
                        clamd_host="127.0.0.1",
                        clamd_port=unused_port,
                        timeout_seconds=0.5,
                    )
                )
            )
            with pytest.raises(ScanUnavailable):
                await ClamAvScanner(cfg).scan_async(b"synthetic scan with no daemon")
        finally:
            await store.close_async()
            await db.close_async()

    asyncio.run(exercise())


def test_real_dependency_outage_retries_then_recovers():
    """Real Postgres state and S3 bytes survive scanner loss and safe retry."""
    _enabled()
    import hashlib
    import socket
    from types import SimpleNamespace

    from src.infrastructure.providers.storage.quarantine_storage import StorageUnavailable

    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    repo = container.resolve(IngestionRepository)
    store = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def exercise():
        await db.start_async()
        await store.start_async()
        try:
            owner = "synthetic-" + uuid4().hex
            source = await repo.create_source_async(owner, "internal", "research")
            content = ("safe-after-network-outage-" + uuid4().hex).encode()
            checksum = hashlib.sha256(content).hexdigest()
            idempotency_key = uuid4().hex
            job, version, new = await repo.stage_async(
                source.id,
                owner,
                "internal",
                "recovery.txt",
                "text/plain",
                len(content),
                checksum,
                idempotency_key,
                3,
            )
            assert new

            # Interrupted upload: a committed intent remains and can be resumed.
            await store.close_async()
            with pytest.raises(StorageUnavailable):
                await store.put_async(version.object_key, content, checksum)
            await repo.fail_staging_async(job.id, owner, "internal", "storage_unavailable")
            failed = await repo.get_job_owned_async(job.id, owner, "internal")
            assert failed[0].status == "dead"
            assert failed[1].state == "failed"
            rearmed, same_version, new = await repo.stage_async(
                source.id,
                owner,
                "internal",
                "recovery.txt",
                "text/plain",
                len(content),
                checksum,
                idempotency_key,
                3,
            )
            assert new and rearmed.id == job.id and same_version.id == version.id
            await store.start_async()
            await store.put_async(same_version.object_key, content, checksum)
            assert await repo.mark_queued_async(job.id, owner, "internal")

            # Scanner endpoint loss is a genuine failed TCP connection, not an
            # optimistic mocked clean verdict. The job must be retry_wait.
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
                probe.bind(("127.0.0.1", 0))
                unreachable = probe.getsockname()[1]
            cfg = SimpleNamespace(
                settings=SimpleNamespace(
                    ingestion=SimpleNamespace(
                        enabled=True,
                        max_file_bytes=1048576,
                        clamd_host="127.0.0.1",
                        clamd_port=unreachable,
                        timeout_seconds=0.5,
                    )
                )
            )
            worker._scanner = ClamAvScanner(cfg)
            assert await worker.run_once_async("scanner-outage")
            waiting = await repo.get_job_owned_async(job.id, owner, "internal")
            assert waiting[0].status == "retry_wait"
            assert waiting[0].error_code == "dependency_unavailable"
            assert waiting[0].attempts == 1
            assert waiting[1].state != "awaiting_approval"

            # The real ClamAV service is healthy again. Move the synthetic
            # retry deadline forward without sleeping through backoff.
            worker._scanner = container.resolve(ClamAvScanner)
            async with db.transaction_async() as session:
                row = await session.get(IngestionJob, job.id)
                row.next_retry_at = datetime.now(UTC) - timedelta(seconds=1)
            assert await worker.run_once_async("scanner-restored")
            result = await repo.get_job_owned_async(job.id, owner, "internal")
            assert result[0].status == "completed"
            assert result[0].attempts == 2
            assert result[1].state == "awaiting_approval"
        finally:
            await store.close_async()
            await db.close_async()

    asyncio.run(exercise())


def test_killed_independent_worker_reclaims_expired_real_db_lease():
    """Kill a real worker after SQL lease commit, then resume via actual clamd."""
    _enabled()
    import hashlib

    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    repo = container.resolve(IngestionRepository)
    store = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def exercise():
        await db.start_async()
        await store.start_async()
        gate = asyncio.Event()
        accepted = asyncio.Event()

        async def blocked_scanner(reader, writer):
            accepted.set()
            await gate.wait()
            writer.close()
            await writer.wait_closed()

        listener = await asyncio.start_server(blocked_scanner, "127.0.0.1", 0)
        process = None
        try:
            actor = "synthetic-" + uuid4().hex
            source = await repo.create_source_async(actor, "internal", "research")
            content = ("crash-lease-" + uuid4().hex).encode()
            checksum = hashlib.sha256(content).hexdigest()
            job, version, created = await repo.stage_async(
                source.id,
                actor,
                "internal",
                "workercrash.txt",
                "text/plain",
                len(content),
                checksum,
                uuid4().hex,
                2,
            )
            assert created
            await store.put_async(version.object_key, content, checksum)
            assert await repo.mark_queued_async(job.id, actor, "internal")
            env = os.environ.copy()
            env["SCB_INGESTION_CLAMD_PORT"] = str(listener.sockets[0].getsockname()[1])
            process = subprocess.Popen(
                [sys.executable, "-m", "src.host.ingestion_worker"],
                env=env,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            await asyncio.wait_for(accepted.wait(), timeout=15)
            locked = await repo.get_job_owned_async(job.id, actor, "internal")
            assert locked[0].status == "leased"
            assert locked[0].attempts == 1
            stale_token = locked[0].lease_owner
            process.kill()
            await asyncio.to_thread(process.wait, 10)
            process = None
            assert (await repo.get_job_owned_async(job.id, actor, "internal"))[0].status == "leased"
            # Advance the already committed synthetic lease instead of sleeping 45s.
            async with db.transaction_async() as session:
                locked_row = await session.get(IngestionJob, job.id)
                locked_row.lease_expires_at = datetime.now(UTC) - timedelta(seconds=2)
            assert await worker.run_once_async("post-crash-recovery")
            assert not await repo.finish_async(job.id, stale_token, True)
            recovered = await repo.get_job_owned_async(job.id, actor, "internal")
            assert recovered[0].status == "completed"
            assert recovered[0].attempts == 2
            assert recovered[1].state == "awaiting_approval"
        finally:
            if process is not None and process.poll() is None:
                process.kill()
                await asyncio.to_thread(process.wait, 10)
            gate.set()
            listener.close()
            await listener.wait_closed()
            await store.close_async()
            await db.close_async()

    asyncio.run(exercise())


def test_real_digital_pdf_upload_scan_and_approval_wait(client):
    """Digital PDF is staged and scanned by real services without publishing."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    http, token, _ = client
    owner = "synthetic-" + uuid4().hex
    bearer = {"Authorization": "Bearer " + token(owner, ["knowledge:test:ingest"])}
    source_response = http.post("/api/v1/knowledge/sources", headers=bearer, json={"domain": "research"})
    assert source_response.status_code == 201
    source_id = source_response.json()["id"]
    pdf = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=180, height=180)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 12 100 Td (Synthetic digital PDF evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(pdf)
    payload = pdf.getvalue()
    uploaded = http.post(
        f"/api/v1/knowledge/sources/{source_id}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex},
        files={"file": ("synthetic.pdf", payload, "application/pdf")},
    )
    assert uploaded.status_code == 202, uploaded.text
    job_id = uploaded.json()["id"]
    version_id = uploaded.json()["version_id"]

    db_container = bootstrap_di()
    db = db_container.resolve(PsqlDbContext)
    storage = db_container.resolve(QuarantineStorage)
    worker = db_container.resolve(IngestionWorkerService)

    async def process():
        await db.start_async()
        try:
            await storage.start_async()
            try:
                assert await worker.run_once_async("pdf-real-scanner")
            finally:
                await storage.close_async()
        finally:
            await db.close_async()

    asyncio.run(process())
    status = http.get(f"/api/v1/knowledge/jobs/{job_id}", headers=bearer)
    assert status.status_code == 200
    assert status.json()["status"] == "completed"
    version = http.get(f"/api/v1/knowledge/versions/{version_id}", headers=bearer)
    assert version.status_code == 200
    assert version.json()["state"] == "awaiting_approval"
    assert version.json()["content_type"] == "application/pdf"

    # Phase06 extends Phase05's atomic processing path: no awaiting_approval
    # without durable, true source-mapped PDF text chunks in PostgreSQL.
    from sqlalchemy import select

    from src.infrastructure.models.entities import KnowledgeChunk

    async def assert_real_pdf_chunks():
        db = bootstrap_di().resolve(PsqlDbContext)
        await db.start_async()
        try:
            async with db.transaction_async() as session:
                chunks = (
                    await session.scalars(
                        select(KnowledgeChunk)
                        .where(
                            KnowledgeChunk.version_id == UUID(version_id),
                            KnowledgeChunk.deployment_scope == "internal",
                        )
                        .order_by(KnowledgeChunk.ordinal)
                    )
                ).all()
                assert chunks
                assert all(item.page_number == 1 and item.extractor_method == "digital_pdf" for item in chunks)
                assert "Synthetic digital PDF evidence" in chunks[0].canonical_text
                from src.core.services.persian_text.persian_processing import normalize_persian

                assert normalize_persian(chunks[0].original_text) == chunks[0].canonical_text
                assert all(item.source_end > item.source_start for item in chunks)
                assert all(item.processor_version == "fa-structure-v1" for item in chunks)
                assert all(item.normalizer_version == "fa-nfc-normalizer-v1" for item in chunks)
                assert all(item.engine_version.startswith("pypdf/") for item in chunks)
        finally:
            await db.close_async()

    asyncio.run(assert_real_pdf_chunks())


@pytest.mark.parametrize("kind", ["docx", "scanned_pdf"])
def test_real_phase06_source_storage_scanner_worker_and_chunk_spans(client, kind):
    """Actual synthetic API/S3/ClamAV/PostgreSQL and OCR or DOCX parser E2E."""
    import hashlib
    import io

    from sqlalchemy import select

    from src.core.services.persian_text.persian_processing import chunk_pages, normalize_persian
    from src.infrastructure.models.entities import KnowledgeChunk
    from src.infrastructure.providers.document.document_extractor import DocumentExtractor

    http, token, _ = client
    owner = "synthetic-" + uuid4().hex
    bearer = {"Authorization": "Bearer " + token(owner, ["knowledge:test:ingest"])}
    response = http.post("/api/v1/knowledge/sources", headers=bearer, json={"domain": "research"})
    assert response.status_code == 201
    source_id = response.json()["id"]

    if kind == "docx":
        from docx import Document

        document = Document()
        document.add_heading("بخش پژوهش", level=1)
        document.add_paragraph("سؤال: چه اقدامی لازم است؟")
        document.add_paragraph("پاسخ: تایید کارشناس ضروری است.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "عنوان", "پاسخ"
        table.cell(1, 0).text, table.cell(1, 1).text = "آیین‌نامه", "تاییدشده"
        stream = io.BytesIO()
        document.save(stream)
        filename = "synthetic.docx"
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        data = stream.getvalue()
    else:
        from PIL import Image, ImageDraw, ImageFont

        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 75)
        image = Image.new("RGB", (1300, 240), "white")
        ImageDraw.Draw(image).text((1240, 48), "سلام دنیا", font=font, fill="black", anchor="ra", direction="rtl")
        output = io.BytesIO()
        image.save(output, format="PDF", resolution=180)
        filename, mime, data = "synthetic-scanned.pdf", "application/pdf", output.getvalue()

    uploaded = http.post(
        f"/api/v1/knowledge/sources/{source_id}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex},
        files={"file": (filename, data, mime)},
    )
    assert uploaded.status_code == 202, uploaded.text
    version_id = UUID(uploaded.json()["version_id"])
    job_id = uploaded.json()["id"]

    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    store = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def run_actual_worker_and_verify():
        await db.start_async()
        await store.start_async()
        try:
            assert await worker.run_once_async("phase06-live-" + kind)
            async with db.transaction_async() as session:
                persisted = (
                    await session.scalars(
                        select(KnowledgeChunk)
                        .where(
                            KnowledgeChunk.version_id == version_id,
                            KnowledgeChunk.deployment_scope == "internal",
                        )
                        .order_by(KnowledgeChunk.ordinal)
                    )
                ).all()
                assert persisted
                original_pages = await DocumentExtractor().extract_async(data, mime)
                expected = chunk_pages(version_id, original_pages)
                assert [item.id for item in persisted] == [item.id for item in expected]
                for saved, item in zip(persisted, expected, strict=True):
                    assert saved.original_text == item.original
                    assert saved.canonical_text == normalize_persian(saved.original_text)
                    assert saved.source_start == item.start
                    assert saved.source_end == item.end
                    assert saved.page_number == item.page
                    assert saved.checksum_sha256 == hashlib.sha256(saved.canonical_text.encode()).hexdigest()
                    assert saved.normalizer_version == "fa-nfc-normalizer-v1"
                    assert saved.processor_version == "fa-structure-v1"
                    assert saved.engine_version == item.engine_version
                    original_page = next(page for page in original_pages if page.number == saved.page_number)
                    assert original_page.text[saved.source_start : saved.source_end] == saved.original_text
                text = "\n".join(row.canonical_text for row in persisted)
                if kind == "docx":
                    assert "عنوان | پاسخ" in text
                    assert "تاییدشده" in text
                    assert all(item.page_number is None and item.extractor_method == "docx" for item in persisted)
                    assert all(item.engine_version.startswith("python-docx/") for item in persisted)
                else:
                    assert "سلام دنیا" in text
                    assert all(item.page_number == 1 and item.extractor_method == "tesseract_ocr" for item in persisted)
                    assert all("tesseract" in item.engine_version for item in persisted)
        finally:
            await store.close_async()
            await db.close_async()

    asyncio.run(run_actual_worker_and_verify())
    status = http.get(f"/api/v1/knowledge/jobs/{job_id}", headers=bearer)
    assert status.status_code == 200
    assert status.json()["status"] == "completed"
    version = http.get(f"/api/v1/knowledge/versions/{version_id}", headers=bearer)
    assert version.status_code == 200
    assert version.json()["state"] == "awaiting_approval"
    assert "published" not in version.text


def test_real_source_metadata_limit_and_disable_lifecycle(client):
    """Synthetic policy cannot be escalated; owner can stop future uploads."""
    http, token, app = client
    actor = "synthetic-" + uuid4().hex
    h = {"Authorization": "Bearer " + token(actor, ["knowledge:test:ingest"])}
    bad = http.post("/api/v1/knowledge/sources", headers=h, json={"domain": "research", "classification": "secret"})
    assert bad.status_code == 422
    bad_limit = http.post(
        "/api/v1/knowledge/sources", headers=h, json={"domain": "research", "max_file_bytes": 5242880}
    )
    assert bad_limit.status_code == 422

    created = http.post("/api/v1/knowledge/sources", headers=h, json={"domain": "research", "max_file_bytes": 1024})
    assert created.status_code == 201, created.text
    source = created.json()
    assert source["classification"] == "synthetic_unclassified"
    assert source["policy_ref"] == "synthetic-only"
    assert source["max_file_bytes"] == 1024

    size_denied = http.post(
        f"/api/v1/knowledge/sources/{source['id']}/documents",
        headers={**h, "Idempotency-Key": uuid4().hex},
        files={"file": ("large.txt", b"a" * 1025, "text/plain")},
    )
    assert size_denied.status_code == 413

    outsider = {"Authorization": "Bearer " + token("synthetic-outsider", ["knowledge:test:ingest"])}
    foreign = http.post(f"/api/v1/knowledge/sources/{source['id']}/disable", headers=outsider)
    assert foreign.status_code == 404
    stopped = http.post(f"/api/v1/knowledge/sources/{source['id']}/disable", headers=h)
    assert stopped.status_code == 200
    assert stopped.json()["status"] == "disabled"
    assert http.post(f"/api/v1/knowledge/sources/{source['id']}/disable", headers=h).status_code == 200
    rejected = http.post(
        f"/api/v1/knowledge/sources/{source['id']}/documents",
        headers={**h, "Idempotency-Key": uuid4().hex},
        files={"file": ("ok.txt", b"synthetic clean bytes", "text/plain")},
    )
    assert rejected.status_code == 403
    assert http.get(f"/api/v1/knowledge/sources/{source['id']}", headers=h).json()["status"] == "disabled"


def test_previously_accepted_job_completes_after_source_disable(client):
    """Phase05 disable denies new uploads, not already-accepted synthetic jobs."""
    http, token, _ = client
    owner = "synthetic-" + uuid4().hex
    bearer = {"Authorization": "Bearer " + token(owner, ["knowledge:test:ingest"])}
    created = http.post("/api/v1/knowledge/sources", headers=bearer, json={"domain": "research"})
    assert created.status_code == 201
    source_id = created.json()["id"]
    uploaded = http.post(
        f"/api/v1/knowledge/sources/{source_id}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex},
        files={"file": ("accepted.txt", b"Synthetic accepted document", "text/plain")},
    )
    assert uploaded.status_code == 202, uploaded.text
    job_id, version_id = uploaded.json()["id"], uploaded.json()["version_id"]

    stopped = http.post(f"/api/v1/knowledge/sources/{source_id}/disable", headers=bearer)
    assert stopped.status_code == 200
    assert stopped.json()["status"] == "disabled"
    denied = http.post(
        f"/api/v1/knowledge/sources/{source_id}/documents",
        headers={**bearer, "Idempotency-Key": uuid4().hex},
        files={"file": ("new.txt", b"new synthetic document", "text/plain")},
    )
    assert denied.status_code == 403

    container = bootstrap_di()
    db = container.resolve(PsqlDbContext)
    storage = container.resolve(QuarantineStorage)
    worker = container.resolve(IngestionWorkerService)

    async def process_existing_job():
        await db.start_async()
        try:
            await storage.start_async()
            try:
                assert await worker.run_once_async("source-disabled-existing-job")
            finally:
                await storage.close_async()
        finally:
            await db.close_async()

    asyncio.run(process_existing_job())
    job = http.get(f"/api/v1/knowledge/jobs/{job_id}", headers=bearer)
    assert job.status_code == 200
    assert job.json()["status"] == "completed"
    version = http.get(f"/api/v1/knowledge/versions/{version_id}", headers=bearer)
    assert version.status_code == 200
    assert version.json()["state"] == "awaiting_approval"
    assert http.get(f"/api/v1/knowledge/sources/{source_id}", headers=bearer).json()["status"] == "disabled"
